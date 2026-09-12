#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""export_docx_gse31210.py — 把 manuscript_gse31210.md 转成 docx 供编辑"""
import sys, os, argparse
sys.path.insert(0, os.path.dirname(__file__))

def md_to_docx(md_path, docx_path):
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    
    with open(md_path, encoding="utf-8") as f:
        lines = f.readlines()
    
    i = 0
    while i < len(lines):
        line = lines[i].rstrip("\n")
        
        # Headers
        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=0)
        elif line.startswith("## "):
            p = doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            p = doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("#### "):
            p = doc.add_heading(line[5:].strip(), level=3)
        # Horizontal rule
        elif line.strip() == "---":
            doc.add_paragraph("─" * 40)
        # Empty line
        elif not line.strip():
            pass
        # Bullet list
        elif line.startswith("- "):
            p = doc.add_paragraph(line[2:], style="List Bullet")
        # Bold text inline
        else:
            # Handle **bold** markers
            parts = []
            current = ""
            j = 0
            while j < len(line):
                if line[j:j+2] == "**":
                    if current:
                        parts.append(("text", current))
                        current = ""
                    j += 2
                    bold_text = ""
                    while j < len(line) and line[j:j+2] != "**":
                        bold_text += line[j]
                        j += 1
                    parts.append(("bold", bold_text))
                    j += 2
                else:
                    current += line[j]
                    j += 1
            if current:
                parts.append(("text", current))
            
            p = doc.add_paragraph()
            for kind, text in parts:
                run = p.add_run(text)
                if kind == "bold":
                    run.bold = True
        
        i += 1
    
    doc.save(docx_path)
    print(f"Docx saved to {docx_path}")

if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--md", required=True)
    parser.add_argument("--docx", required=True)
    args = parser.parse_args()
    md_to_docx(args.md, args.docx)
